import io

from openpyxl import load_workbook


def extract_excel_text(file_bytes: bytes) -> str:
    """Extract cell text from every sheet in an .xlsx file, one row per line.

    data_only=True reads formulas' last-cached computed value instead of the
    formula string. Excel has no page concept, so callers store page=None
    for these chunks, same as .docx.
    """
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)

    parts = []
    for sheet in workbook.worksheets:
        sheet_rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell).strip() for cell in row if cell is not None and str(cell).strip()]
            if cells:
                sheet_rows.append(" | ".join(cells))

        if sheet_rows:
            parts.append(f"## Sheet: {sheet.title}")
            parts.extend(sheet_rows)

    return "\n\n".join(parts)

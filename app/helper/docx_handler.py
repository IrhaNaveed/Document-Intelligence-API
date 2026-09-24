import io

from docx import Document as DocxDocument


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraph and table text from a .docx file.

     Caller stores page=None for these chunks instead of guessing one.
    """
    document = DocxDocument(io.BytesIO(file_bytes))

    parts = [para.text for para in document.paragraphs if para.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n\n".join(parts)
